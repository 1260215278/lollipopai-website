# -*- coding: utf-8 -*-
import openpyxl
from pathlib import Path

SRC = Path(r"C:\Users\Administrator\Downloads\https___www.lollipop.im_-Coverage-2026-10-13.xlsx")
OUT = Path(r"C:\Users\Administrator\WorkBuddy\Worktrees\lollipop\master-17dd9e0b\coverage_dump.txt")

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
lines = []
lines.append(f"FILE: {SRC.name}")
lines.append(f"SHEETS: {wb.sheetnames}")
lines.append("=" * 80)

for ws in wb.worksheets:
    lines.append("")
    lines.append(f"### SHEET: {ws.title}  (max_row={ws.max_row}, max_column={ws.max_column})")
    lines.append("-" * 80)
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        # skip fully-empty rows
        if all(c is None for c in row):
            continue
        cells = ["" if c is None else str(c) for c in row]
        lines.append(f"R{r}: " + "\t".join(cells))

OUT.write_text("\n".join(lines), encoding="utf-8", newline="")
print(f"Dumped {len(wb.sheetnames)} sheet(s) to {OUT}")
for ws in wb.worksheets:
    print(f"  - {ws.title}: {ws.max_row} rows")
