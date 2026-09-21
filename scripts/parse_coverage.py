import openpyxl, warnings
warnings.filterwarnings("ignore")

fp = r"C:\Users\Administrator\Downloads\https___www.lollipop.im_-Coverage-2026-09-17.xlsx"
wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)

# Chart sheet tail
ws = wb["图表"]
rows = list(ws.iter_rows(values_only=True))
print("=== 图表 HEAD ===")
for r in rows[:3]:
    print(r)
print("=== 图表 TAIL (latest 20) ===")
for r in rows[-20:]:
    print(r)

# Severe issues totals
ws2 = wb["严重问题"]
rows2 = list(ws2.iter_rows(values_only=True))
print("\n=== 严重问题 (full) ===")
for r in rows2:
    print(r)
failed = 0
for r in rows2[1:]:
    reason, src, verify, cnt = r
    try:
        c = float(cnt)
    except (TypeError, ValueError):
        c = 0
    if verify == "失败":
        failed += c
print(f"\nSUM of 验证=失败: {failed:.0f}")

wb.close()
